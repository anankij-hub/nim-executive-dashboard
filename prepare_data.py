from pathlib import Path
import csv
import json

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "data"
OUTPUT_DIR.mkdir(exist_ok=True)
INPUT_DIR.mkdir(exist_ok=True)

SUPPORTED = {".xlsx", ".xls", ".csv"}
files = sorted([p for p in INPUT_DIR.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED], key=lambda p: p.name.lower())

if not files:
    print("ไม่พบไฟล์ Excel/CSV ในโฟลเดอร์ input/")
    print(f"ให้อัปโหลดไฟล์ผ่านหน้าเว็บ หรือวางไฟล์ที่: {INPUT_DIR}")
    raise SystemExit(1)

inventory = []
errors = []


def clean_headers(values):
    out = []
    for i, v in enumerate(values, start=1):
        text = "" if v is None else str(v).strip()
        out.append(text or f"Column_{i}")
    return out


def read_xlsx(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("ต้องติดตั้ง openpyxl: python3 -m pip install openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            columns = clean_headers(first)
            rows = max((ws.max_row or 1) - 1, 0)
            sheets.append({"sheet": ws.title, "rows": int(rows), "columns": columns})
    finally:
        wb.close()
    return sheets


def read_xls(path: Path):
    try:
        import pandas as pd
    except ImportError as e:
        raise RuntimeError("ไฟล์ .xls ต้องติดตั้ง pandas และ xlrd: python3 -m pip install pandas xlrd") from e
    try:
        xls = pd.ExcelFile(path)
    except Exception as e:
        raise RuntimeError(f"เปิด .xls ไม่ได้ ({e}) — ลองติดตั้ง xlrd") from e
    sheets = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet)
        sheets.append({"sheet": str(sheet), "rows": int(len(df)), "columns": clean_headers(df.columns.tolist())})
    return sheets


def read_csv(path: Path):
    encodings = ["utf-8-sig", "utf-8", "cp874", "tis-620"]
    last_err = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                header = next(reader, [])
                rows = sum(1 for _ in reader)
            return [{"sheet": "CSV", "rows": rows, "columns": clean_headers(header), "encoding": enc}]
        except UnicodeDecodeError as e:
            last_err = e
    raise RuntimeError(f"อ่าน encoding ของ CSV ไม่ได้: {last_err}")


print("\n=== NIM DATA INVENTORY ===\n")
for path in files:
    print(f"ไฟล์: {path.name}")
    try:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            sheets = read_xlsx(path)
        elif suffix == ".xls":
            sheets = read_xls(path)
        else:
            sheets = read_csv(path)

        info = {"file": path.name, "size_bytes": path.stat().st_size, "sheets": sheets}
        inventory.append(info)
        for sh in sheets:
            print(f"  Sheet: {sh['sheet']}")
            print(f"    จำนวนแถว: {sh['rows']:,}")
            print("    Columns:")
            for c in sh["columns"]:
                print(f"      - {c}")
    except Exception as e:
        errors.append({"file": path.name, "error": str(e)})
        print(f"  อ่านไฟล์ไม่ได้: {e}")
    print()

output = {"generated_from": "local_upload", "files": inventory, "errors": errors}
# Keep backwards-compatible top-level list because the browser currently expects a list.
output_path = OUTPUT_DIR / "data_inventory.json"
with output_path.open("w", encoding="utf-8") as f:
    json.dump(inventory, f, ensure_ascii=False, indent=2)

with (OUTPUT_DIR / "data_inventory_report.json").open("w", encoding="utf-8") as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print("เสร็จแล้ว")
print(f"บันทึกโครงสร้างข้อมูลไว้ที่: {output_path}")
if errors:
    print(f"มีไฟล์ที่อ่านไม่สำเร็จ {len(errors)} ไฟล์ — ดูรายละเอียดใน data/data_inventory_report.json")
